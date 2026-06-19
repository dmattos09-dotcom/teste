from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('BRIEFING — LANDING PAGE DE EVENTO PRESENCIAL')
run.bold = True
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = sub.add_run('Kiwify + Domínio Próprio')
run2.font.size = Pt(12)
run2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
run2.italic = True

doc.add_paragraph()

def add_section(doc, number, title_text):
    p = doc.add_paragraph()
    run = p.add_run(f'{number}. {title_text}')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    run.underline = True

def add_question(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(11)

sections_data = [
    ("1", "INFORMAÇÕES DO EVENTO", [
        "Qual é o nome do evento?",
        "Qual é o tema/nicho do evento? (ex: marketing, saúde, desenvolvimento pessoal, negócios, etc.)",
        "Qual é a data (dia, mês, ano) e o horário de início e encerramento?",
        "É um único dia ou múltiplos dias? Se múltiplos, quais as datas?",
        "Qual é o local (nome do espaço/venue, endereço completo, cidade e estado)?",
        "Qual é a capacidade máxima de participantes?",
        "Haverá transmissão online (híbrido) ou é 100% presencial?",
        "Qual é a transformação/resultado que o participante vai ter ao final do evento? (o “depois” do antes x depois)",
    ]),
    ("2", "PÚBLICO-ALVO", [
        "Quem é o público ideal do evento? (profissão, faixa etária, interesses)",
        "Quais são as principais dores que esse público tem?",
        "Quais objeções esse público costuma ter na hora de comprar um ingresso?",
        "O público já conhece o cliente/marca, ou é uma audiência fria?",
    ]),
    ("3", "CLIENTE / MARCA", [
        "Qual é o nome do cliente (pessoa física ou empresa)?",
        "Qual é o nome da marca (se diferente)?",
        "Tem logo? Se sim, enviar em PNG com fundo transparente (alta resolução).",
        "Quais são as cores da marca (códigos hex, se possível)?",
        "Existe uma identidade visual já definida para o evento? (paleta de cores, fontes, arte do evento)",
        "Quais são os perfis nas redes sociais do cliente/evento? (Instagram, YouTube, etc.)",
        "O cliente tem site ou outras páginas ativas?",
    ]),
    ("4", "PALESTRANTES / ATRAÇÕES", [
        "Quem são os palestrantes ou atrações do evento?",
        "Para cada um: nome completo, foto profissional, mini bio e tema da palestra/workshop.",
        "O cliente/host também vai palestrar? Se sim, incluir os dados dele também.",
        "Haverá palestrantes surpresa ou que não podem ser divulgados? Como tratar isso na página?",
    ]),
    ("5", "PROGRAMAÇÃO", [
        "Existe uma grade/programação definida? (horários, blocos, atividades)",
        "Quais são os tópicos e módulos que serão abordados no evento?",
        "Haverá coffee break, almoço, material físico, brindes incluído no ingresso?",
        "Terá certificado de participação?",
    ]),
    ("6", "INGRESSOS E PREÇOS", [
        "Quantos tipos de ingresso existem? (ex: pista, VIP, premium)",
        "Para cada tipo: nome, preço, o que está incluído e o que diferencia.",
        "Há venda em lotes (lote 1, lote 2…)? Se sim, quais os preços e prazos de cada lote?",
        "Qual é o prazo final de vendas / encerramento das inscrições?",
        "Quais são as formas de pagamento disponíveis na Kiwify? (cartão, PIX, boleto, parcelamento)",
        "Em quantas parcelas pode parcelar? Tem juros?",
        "Já tem o link do produto na Kiwify ou ainda será criado?",
    ]),
    ("7", "PROVA SOCIAL", [
        "Já houve edições anteriores do evento? Se sim, quantas pessoas participaram?",
        "Tem fotos ou vídeos de edições anteriores para usar na página?",
        "Tem depoimentos de participantes anteriores? (texto, áudio ou vídeo)",
        "Existem logos de parceiros, patrocinadores ou apoiadores para exibir?",
        "O cliente tem números de autoridade para mostrar? (ex: X alunos, X anos de mercado, X seguidores)",
    ]),
    ("8", "TEXTOS E COPY", [
        "Tem alguma headline (frase principal) já pensada, ou deixa a cargo da criação?",
        "Tem texto de descrição do evento pronto, ou precisa ser criado do zero?",
        "Quais são as principais perguntas frequentes (FAQ) que os participantes costumam fazer?",
        "Existe algum tom de voz da marca? (ex: descontraído, técnico, inspiracional, direto)",
        "Tem alguma referência de copy que o cliente gosta?",
    ]),
    ("9", "REFERÊNCIAS VISUAIS", [
        "Tem referências de landing pages que o cliente gosta (pode ser de outros eventos ou segmentos)?",
        "Tem preferência de estilo visual? (ex: clean/minimalista, dark, colorido, luxuoso, moderno)",
        "Quais imagens e materiais gráficos estão disponíveis? (fotos do cliente, do local, banners, etc.)",
        "Tem um banner ou arte oficial do evento já criado?",
    ]),
    ("10", "CONFIGURAÇÕES TÉCNICAS", [
        "Qual é o domínio que vai ser usado? (ex: evento.nomedomarca.com.br)",
        "O domínio já está apontado para a Kiwify ou isso ainda precisa ser configurado?",
        "Terá pixel do Facebook/Meta para rastreamento? Se sim, qual o ID do pixel?",
        "Terá Google Analytics ou Google Tag Manager? Se sim, qual o ID?",
        "Terá remarketing ou integração com alguma ferramenta de e-mail marketing?",
        "Haverá botão de WhatsApp na página para dúvidas? Se sim, qual o número?",
    ]),
    ("11", "INFORMAÇÕES LEGAIS", [
        "Qual é o CNPJ ou CPF do responsável pela venda (para rodapé da página)?",
        "Qual é a política de reembolso do evento?",
        "Precisa de termos de uso ou política de privacidade na página?",
    ]),
    ("12", "PRAZOS", [
        "Qual é a data limite para a landing page estar no ar?",
        "Haverá rodadas de revisão? Quantas?",
        "Quem é o responsável por aprovar a página (o próprio cliente ou outra pessoa)?",
    ]),
]

for num, title_text, questions in sections_data:
    add_section(doc, num, title_text)
    for q in questions:
        add_question(doc, q)
    doc.add_paragraph()

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(
    'Preencha as respostas de forma corrida — não precisa seguir o formato acima.\n'
    'Com todas as informações em mãos, a landing page será construída completa.'
)
run.italic = True
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

doc.save('/home/user/teste/Briefing_Landing_Page_Evento.docx')
print("Arquivo criado com sucesso!")
